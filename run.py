#!/usr/bin/env python3

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm,Pt
import glob
import sys

symbols={'出勤': '√',
'到岗': '√',
 '公差': '○',
 '出差': '○',
 '补休': '□',
 '年休假': '△',
 '年假': '△',
 '事假': 'S',
 '病假': 'B',
 '婚假': 'H',
 '丧假': 'SJ',
 '探亲': 'T',
 '产假': 'CJ',
 '育儿假': 'Y',
 '护理假': 'HL',
 '工伤': 'G',
 '迟到': 'C',
 '早退': 'Z',
 '旷工': 'K',
 '周末': '',
 '放假': '',
 }

yyyymm=sys.argv[1]
yyyy=yyyymm[:4]
mm=yyyymm[-2:]
irow=(5,6)
template='template.docx'
document = Document(template)

document.paragraphs[2].text=f'部门： 某部门                                                                                                    {yyyy}年 {mm} 月'
document.paragraphs[2].runs[0].font.name='仿宋_GB2312'
document.paragraphs[2].runs[0].font.size=152400

    
table=document.tables[0]
for f in sorted(glob.glob(f'{yyyymm}??.md')):
    dd=int(f[6:8])
    for i, s in zip(irow, [symbols[d] for d in open(f).readline().strip().split()]):
        c=table.columns[dd+1].cells[i]
        c.text=s
        c.paragraphs[0].runs[0].font.name='仿宋_GB2312'
        c.paragraphs[0].runs[0].font.size='152400'
        c.paragraphs[0].alignment=WD_PARAGRAPH_ALIGNMENT.CENTER

document.save(f'考勤表_{yyyymm}_张三.docx')

